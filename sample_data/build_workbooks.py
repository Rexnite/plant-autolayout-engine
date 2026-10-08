"""Regenerate generalized example workbooks from their JSON sources."""

import json
from pathlib import Path

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parent
TABLES = (("Boundary", "boundary"), ("Groups", "system_groups"), ("Units", "units"),
          ("Flow", "flow_edges"), ("Proximity", "proximity_relationships"),
          ("Constraints", "constraints"))


def build(source: Path) -> Path:
    project = json.loads(source.read_text(encoding="utf-8"))
    workbook = Workbook()
    workbook.remove(workbook.active)
    info = workbook.create_sheet("Project")
    info.append(["key", "value"])
    for key in ("plant_type", "layout_id"):
        info.append([key, project[key]])
    for key, value in project.get("metadata", {}).items():
        info.append([key, value])
    for sheet_name, section in TABLES:
        rows = project.get(section, [])
        if not rows:
            continue
        sheet = workbook.create_sheet(sheet_name)
        columns = list(dict.fromkeys(column for row in rows for column in row))
        sheet.append(columns)
        for row in rows:
            sheet.append([json.dumps(row.get(column)) if isinstance(row.get(column), (dict, list))
                          else row.get(column) for column in columns])
    target = source.with_suffix(".xlsx")
    workbook.save(target)
    return target


if __name__ == "__main__":
    for name in ("sample_general_wtp_input.json", "sample_petrochemical_input.json"):
        print(build(ROOT / name))
