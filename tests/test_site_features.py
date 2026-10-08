"""Typed site input and site-aware pipeline regressions."""

import copy
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from wtp_autolayout.candidates import CandidateGenerator
from wtp_autolayout.constraints.engine import default_core_engine
from wtp_autolayout.constraints.site import access_results, group_relationship_results, site_results
from wtp_autolayout.export import JSONExporter, PDFExporter
from wtp_autolayout.geometry.routing import FlowRouter, RoutingStatus
from wtp_autolayout.geometry import GeometryEngine
from wtp_autolayout.geometry.site import corridor_geometry, hard_site_obstacles
from wtp_autolayout.models import (ConstraintClassification, Coordinate, GroupRelationship, GroupRelationshipType,
                                   Layout, Placement, SiteZone)
from wtp_autolayout.optimization.local_search import propose_neighbor
from wtp_autolayout.parsers.errors import ParserError
from wtp_autolayout.parsers.json_parser import parse_project_data
from wtp_autolayout.parsers.excel_parser import parse_excel_file
from wtp_autolayout.scoring import ScoringEngine, ScoringError
from wtp_autolayout.scenarios import ScenarioResult, compare_scenarios
from wtp_autolayout.scoring.objectives import EngineeringObjectiveCalculator
from wtp_autolayout.visualization import LayoutRenderer, PlotlyLayoutRenderer


ROOT = Path(__file__).resolve().parents[1]


def square(x0, y0, x1, y1):
    return [{"easting": x0, "northing": y0}, {"easting": x1, "northing": y0},
            {"easting": x1, "northing": y1}, {"easting": x0, "northing": y1}]


def base_data():
    return {"layout_id": "site_test", "plant_type": "petrochemical",
            "boundary": square(0, 0, 100, 100),
            "units": [{"unit_id": "A", "name": "A", "category": "process", "shape_type": "rectangle",
                       "width": 10, "length": 10, "required_clearance": 2},
                      {"unit_id": "B", "name": "B", "category": "process", "shape_type": "rectangle",
                       "width": 10, "length": 10, "required_clearance": 2}],
            "flow_edges": [{"edge_id": "F", "source_unit_id": "A", "target_unit_id": "B",
                            "flow_type": "other"}]}


def placed(data, a=(20, 50), b=(80, 50)):
    layout = parse_project_data(data)
    layout.placements = [Placement("a", "A", Coordinate(*a)), Placement("b", "B", Coordinate(*b))]
    return layout


class SiteFeatureTests(unittest.TestCase):
    def test_zone_validation_and_roundtrip(self):
        data = base_data()
        data["site_zones"] = [{"zone_id": "z", "name": "No build", "zone_type": "exclusion",
                               "vertices": square(40, 40, 60, 60)}]
        layout = parse_project_data(data)
        self.assertEqual(Layout.from_dict(layout.to_dict()).site_zones[0].zone_id, "z")
        with self.assertRaises(ValueError):
            SiteZone("bad", "Bad", [Coordinate(0, 0), Coordinate(1, 0)], "exclusion")
        data["site_zones"][0]["zone_type"] = "unknown"
        with self.assertRaises(ParserError):
            parse_project_data(data)

    def test_hard_exclusion_and_fixed_structure(self):
        data = base_data()
        data["site_zones"] = [{"zone_id": "z", "name": "No build", "zone_type": "exclusion",
                               "classification": "hard", "vertices": square(17, 47, 23, 53)}]
        data["existing_structures"] = [{"structure_id": "fixed", "name": "Fixed",
                                         "vertices": square(77, 47, 83, 53)}]
        layout = placed(data)
        self.assertEqual(len(site_results(layout)), 2)
        self.assertFalse(ScoringEngine(default_core_engine()).score_layout(layout).passed_hard_constraints)
        self.assertEqual(len(hard_site_obstacles(layout)), 2)
        for _ in range(3):
            candidate = CandidateGenerator(random_seed=42).generate(layout, replace_existing=True)
            self.assertFalse(site_results(candidate))

    def test_soft_reserved_and_explicit_expansion(self):
        data = base_data()
        data["site_zones"] = [{"zone_id": "reserve", "name": "Reserve", "zone_type": "reserved",
                               "classification": "soft", "weight": 0.5,
                               "vertices": square(16, 46, 24, 54)},
                              {"zone_id": "exp", "name": "Expansion", "zone_type": "expansion",
                               "classification": "soft", "vertices": square(70, 40, 90, 60)}]
        layout = placed(data)
        rows = site_results(layout)
        self.assertEqual(len(rows), 2)
        score = ScoringEngine(default_core_engine()).score_layout(layout)
        self.assertTrue(score.passed_hard_constraints)
        self.assertGreater(score.soft_penalty, 0)
        zone = EngineeringObjectiveCalculator()._expansion_zone(layout)
        self.assertEqual(zone.bounds, (70.0, 40.0, 90.0, 60.0))

    def test_access_and_legacy_roads(self):
        data = base_data()
        data["units"][0].update({"access_required": True, "max_access_distance": 10,
                                  "access_classification": "hard"})
        data["roads"] = [{"road_id": "old", "path": [{"easting": 0, "northing": 10},
                                                   {"easting": 100, "northing": 10}], "width": 4}]
        layout = placed(data, a=(20, 30))
        row = access_results(layout)[0]
        self.assertEqual(row["nearest_corridor"], "old")
        self.assertFalse(row["satisfied"])
        self.assertFalse(ScoringEngine(default_core_engine()).score_layout(layout).passed_hard_constraints)
        layout.placements[0] = Placement("a", "A", Coordinate(20, 20))
        self.assertTrue(access_results(layout)[0]["satisfied"])
        layout.units[0].access_classification = ConstraintClassification.SOFT
        layout.placements[0] = Placement("a", "A", Coordinate(20, 30))
        self.assertTrue(ScoringEngine(default_core_engine()).score_layout(layout).passed_hard_constraints)

    def test_corridor_geometry_and_routing(self):
        data = base_data()
        data["corridors"] = [{"corridor_id": "pipe", "name": "Pipe", "corridor_type": "routing",
                              "width": 6, "path": [{"easting": 10, "northing": 50},
                                                    {"easting": 90, "northing": 50}]}]
        layout = placed(data)
        corridor = corridor_geometry(layout.corridors[0])
        self.assertGreater(corridor.area, 0)
        route = FlowRouter().route(layout)[0]
        self.assertEqual(route.status, RoutingStatus.SUCCESS)
        self.assertGreater(route.corridor_length, 0)

    def test_maintenance_clearance_affects_spacing(self):
        data = base_data()
        data["units"][0]["maintenance_clearance"] = 20
        layout = placed(data, a=(20, 50), b=(45, 50))
        result = GeometryEngine().validate_layout_geometry(layout)
        self.assertTrue(result.clearance_violations)
        data["site_zones"] = [{"zone_id": "z", "name": "No build", "zone_type": "exclusion",
                               "vertices": square(38, 45, 42, 55)}]
        layout = placed(data, a=(20, 50), b=(80, 50))
        self.assertTrue(any(row["unit_id"] == "A" for row in site_results(layout)))

    def test_group_relationship_nearest_footprints(self):
        data = base_data()
        data["system_groups"] = [{"group_id": "g1", "name": "One"},
                                  {"group_id": "g2", "name": "Two"}]
        data["units"][0]["system_group"] = "g1"
        data["units"][1]["system_group"] = "g2"
        data["group_relationships"] = [{"relationship_id": "g", "group_id_a": "g1", "group_id_b": "g2",
                                        "relationship_type": "min_distance", "distance": 60,
                                        "classification": "soft"}]
        layout = placed(data)
        row = group_relationship_results(layout)[0]
        self.assertEqual(row["actual_distance"], 50)
        self.assertEqual(row["distance_basis"], "nearest_member_footprint_gap")
        self.assertFalse(row["satisfied"])
        layout.group_relationships[0].relationship_type = GroupRelationshipType.PREFERRED_NEAR
        self.assertEqual(GroupRelationship.from_dict(layout.group_relationships[0].to_dict()).group_id_a, "g1")

    def test_neighbor_rejects_hard_site(self):
        import random
        data = base_data()
        data["site_zones"] = [{"zone_id": "all", "name": "Blocked", "zone_type": "exclusion",
                               "vertices": square(10, 10, 90, 90)}]
        layout = placed(data)
        self.assertIsNone(propose_neighbor(layout, GeometryEngine(), random.Random(3)))

    def test_engineering_rule_table(self):
        data = base_data()
        data["site_zones"] = [{"zone_id": "reserve", "name": "Reserve", "zone_type": "reserved",
                               "vertices": square(60, 60, 80, 80)}]
        data["engineering_rules"] = [
            {"rule_id": "pair", "scope": "unit_pair", "subject_a": "A", "subject_b": "B",
             "rule_type": "min_distance", "distance": 20, "classification": "hard"},
            {"rule_id": "road", "scope": "unit_access", "subject_a": "A",
             "rule_type": "max_distance", "distance": 30, "classification": "soft"},
            {"rule_id": "reserve_rule", "scope": "zone_occupancy", "subject_a": "reserve",
             "rule_type": "avoid", "classification": "soft", "weight": 0.3},
        ]
        layout = parse_project_data(data)
        self.assertEqual(layout.proximity_relationships[0].relationship_id, "pair")
        self.assertTrue(layout.units[0].access_required)
        self.assertEqual(layout.site_zones[0].weight, 0.3)
        self.assertEqual(layout.metadata["engineering_rules"][0]["rule_id"], "pair")

    def test_excel_optional_site_sheets(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "site.xlsx"
            workbook = load_workbook(ROOT / "sample_data" / "sample_petrochemical_input.xlsx")
            sheet = workbook.create_sheet("SiteZones")
            sheet.append(["zone_id", "name", "zone_type", "classification", "vertices"])
            sheet.append(["z", "No build", "exclusion", "hard", json.dumps(square(4, 4, 12, 12))])
            corridor = workbook.create_sheet("Corridors")
            corridor.append(["corridor_id", "name", "corridor_type", "width", "path"])
            corridor.append(["c", "Access", "access", 4,
                             json.dumps([{"easting": 20, "northing": 20}, {"easting": 90, "northing": 20}])])
            workbook.save(target)
            layout = parse_excel_file(str(target))
            self.assertEqual(layout.site_zones[0].zone_id, "z")
            self.assertEqual(layout.corridors[0].corridor_id, "c")

    def test_fixed_seed_site_candidate(self):
        data = base_data()
        data["site_zones"] = [{"zone_id": "z", "name": "No build", "zone_type": "exclusion",
                               "vertices": square(40, 40, 60, 60)}]
        layout = parse_project_data(data)
        first = CandidateGenerator(random_seed=19).generate(layout)
        second = CandidateGenerator(random_seed=19).generate(layout)
        self.assertEqual(first.to_dict(), second.to_dict())
        self.assertFalse(site_results(first))

    def test_unknown_objective_rejected(self):
        with self.assertRaises(ScoringError):
            ScoringEngine(reward_weight_overrides={"accsess": 5.0})

    def test_site_aware_optimizer_and_scenarios(self):
        from wtp_autolayout.optimization import Optimizer
        data = base_data()
        data["site_zones"] = [{"zone_id": "z", "name": "No build", "zone_type": "exclusion",
                               "vertices": square(40, 40, 60, 60)}]
        layout = parse_project_data(data)
        scorer = ScoringEngine(default_core_engine())
        optimized = Optimizer(CandidateGenerator(random_seed=9), scorer, random_seed=9,
                              initial_candidate_count=2, max_iterations=5,
                              algorithm="local_search").optimize(layout)
        self.assertIsNotNone(optimized.best_layout)
        self.assertFalse(site_results(optimized.best_layout))
        run = SimpleNamespace(layout=optimized.best_layout,
              optimization_result=SimpleNamespace(best_score_result=optimized.best_score_result),
              timings={"total_runtime": 1.0}, exported_files={}, output_dir="")
        first = ScenarioResult.from_run("site A", data, run, 9, 2, 5, "local_search")
        second = ScenarioResult.from_dict({**first.to_dict(), "name": "site B"})
        self.assertEqual(ScenarioResult.from_dict(first.to_dict()).to_dict(), first.to_dict())
        self.assertIn("Site exclusions", compare_scenarios([first, second])[0])

    def test_streamlit_project_helper(self):
        import app
        data = base_data()
        data["site_zones"] = [{"zone_id": "z", "name": "No build", "zone_type": "exclusion",
                               "vertices": square(40, 40, 60, 60)}]
        data["engineering_rules"] = [{"rule_id": "access", "scope": "unit_access",
                                      "subject_a": "A", "rule_type": "max_distance", "distance": 20}]
        clean = app._project_from_tables(data)
        self.assertEqual(parse_project_data(clean).site_zones[0].zone_id, "z")
        self.assertTrue(parse_project_data(clean).units[0].access_required)

    def test_json_excel_visual_and_export(self):
        data = base_data()
        data["site_zones"] = [{"zone_id": "z", "name": "Reserve", "zone_type": "reserved",
                               "classification": "soft", "vertices": square(40, 40, 60, 60)}]
        data["existing_structures"] = [{"structure_id": "s", "name": "Old",
                                         "vertices": square(5, 5, 12, 12)}]
        data["corridors"] = [{"corridor_id": "access", "name": "Access", "corridor_type": "access",
                              "width": 4, "path": [{"easting": 10, "northing": 15},
                                                    {"easting": 90, "northing": 15}]}]
        layout = placed(data)
        before = copy.deepcopy(layout.to_dict())
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            JSONExporter().export_layout(layout, folder / "nested" / "site.json")
            exported = json.loads((folder / "nested" / "site.json").read_text())
            self.assertEqual(exported["existing_structures"][0]["structure_id"], "s")
            figure = PlotlyLayoutRenderer().render(layout)
            self.assertTrue(any(trace.name == "Existing structure" for trace in figure.data))
            png = LayoutRenderer().save_png(layout, folder / "site.png")
            self.assertGreater(Path(png).stat().st_size, 100)
            pdf = PDFExporter().export_layout_report(layout, folder / "site.pdf", include_image=False)
            self.assertGreater(Path(pdf).stat().st_size, 100)
        self.assertEqual(layout.to_dict(), before)


if __name__ == "__main__":
    unittest.main()
